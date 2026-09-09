"""
L'export du journal des ventes.

┌──────────────────────────────────────────────────────────────────────────────┐
│ UN EXPORT PORTE LE PÉRIMÈTRE FILTRÉ, JAMAIS LA PAGE AFFICHÉE.                │
│                                                                              │
│ C'est le défaut que le back-office a dû corriger sur ses niveaux de stock :   │
│ l'écran paginait et filtrait en mémoire, si bien que le fichier téléchargé ne │
│ couvrait pas le même périmètre que l'écran qui l'avait déclenché. Le socle    │
│ `ExportableListMixin` le garantit par construction (`filter_queryset` sans    │
│ pagination) ; ce test le vérifie plutôt que de le croire.                     │
└──────────────────────────────────────────────────────────────────────────────┘

Les tests s'authentifient en GÉRANT et non en propriétaire : `WarehouseScoped`
sort en amont pour un propriétaire, et la moitié du code de périmètre ne serait
pas exécutée. C'est la règle posée par `test_pull_scope_resolves`.
"""
import io
from decimal import Decimal

from django.utils import timezone
from rest_framework import status
from rest_framework.test import APITestCase

from apps.sales.models import Sale
from apps.sales.reports import build_sales_report
from apps.sales.tests._helpers import make_org_with_users
from apps.settings.models import Currency, OrganizationCurrency


class _SetupVentes(APITestCase):
    def setUp(self):
        ctx = make_org_with_users()
        self.__dict__.update(ctx)
        self.client.force_authenticate(user=self.manager)

    @property
    def _headers(self):
        return {'HTTP_X_ORGANIZATION_ID': str(self.org.id)}

    def _vente(self, reference, total, statut=Sale.Status.COMPLETED,
               devise=None, du='0.00'):
        return Sale.objects.create(
            organization=self.org, warehouse=self.warehouse,
            reference=reference, status=statut,
            subtotal=Decimal(total), total=Decimal(total),
            amount_paid=Decimal(total) - Decimal(du), amount_due=Decimal(du),
            currency=devise or self.org.currency,
            sold_by=self.manager, sale_date=timezone.now(),
        )

    def _seconde_devise(self):
        usd, _ = Currency.objects.get_or_create(
            code='USD',
            defaults={'name': 'Dollar', 'symbol': '$', 'decimal_places': 2},
        )
        OrganizationCurrency.objects.create(
            organization=self.org, currency=usd, is_primary=False,
            exchange_rate=Decimal('2800.000000'), is_active=True,
        )


class ContratHttpTests(_SetupVentes):
    """Le contrat de l'endpoint, identique à celui des exports de stock."""

    def test_export_pdf(self):
        self._vente('VT-1', '100.00')
        resp = self.client.get(
            '/api/v1/sales/export/?export_format=pdf', **self._headers,
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertIn('historique_des_ventes', resp['Content-Disposition'])
        self.assertTrue(resp.content.startswith(b'%PDF'))

    def test_export_xlsx(self):
        self._vente('VT-1', '100.00')
        resp = self.client.get(
            '/api/v1/sales/export/?export_format=xlsx', **self._headers,
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        self.assertTrue(resp.content.startswith(b'PK'))

    def test_excel_est_un_alias_de_xlsx(self):
        """Le mot du bouton et le mot du serveur ne doivent pas diverger."""
        self._vente('VT-1', '100.00')
        resp = self.client.get(
            '/api/v1/sales/export/?export_format=excel', **self._headers,
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)

    def test_un_format_inconnu_repond_400_et_non_500(self):
        resp = self.client.get(
            '/api/v1/sales/export/?export_format=docx', **self._headers,
        )
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_un_export_VIDE_reste_un_fichier_valide(self):
        """Zéro vente n'est pas une erreur : c'est un rapport qui le dit."""
        resp = self.client.get(
            '/api/v1/sales/export/?export_format=pdf', **self._headers,
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertTrue(resp.content.startswith(b'%PDF'))


class PerimetreTests(_SetupVentes):
    """
    Le fichier couvre ce que l'écran montre, ni plus ni moins.

    On passe par le VRAI endpoint et on compte les lignes du classeur : c'est
    la seule preuve qui traverse `filter_queryset`, la pagination et le
    périmètre entrepôt. Une fausse requête montée à la main les court-circuite
    tous les trois, et un test qui les court-circuite ne prouve rien.
    """

    def _lignes_du_classeur(self, requete=''):
        from openpyxl import load_workbook

        resp = self.client.get(
            f'/api/v1/sales/export/?export_format=xlsx{requete}', **self._headers,
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK, resp.content[:400])
        feuille = load_workbook(io.BytesIO(resp.content)).active
        # La colonne des références : on relève tout ce qui ressemble à l'une
        # des nôtres, ce qui saute l'entête d'identité et le cartouche.
        return [
            cellule.value
            for rangee in feuille.iter_rows()
            for cellule in rangee
            if isinstance(cellule.value, str) and cellule.value.startswith('VT-')
        ]

    def test_le_filtre_de_STATUT_est_applique(self):
        self._vente('VT-1', '100.00')
        self._vente('VT-2', '50.00', statut=Sale.Status.CANCELLED)

        self.assertEqual(self._lignes_du_classeur('&status=cancelled'), ['VT-2'])

    def test_la_RECHERCHE_est_appliquee(self):
        self._vente('VT-AAA', '100.00')
        self._vente('VT-BBB', '50.00')

        self.assertEqual(self._lignes_du_classeur('&search=AAA'), ['VT-AAA'])

    def test_l_export_n_est_PAS_pagine(self):
        """
        Cent une ventes, page demandée de vingt : le fichier les porte TOUTES.

        Sans cela le marchand télécharge « l'historique » et reçoit sa première
        page, sans que rien ne le dise. C'est le défaut que le back-office a dû
        corriger sur ses niveaux de stock.
        """
        for i in range(101):
            self._vente(f'VT-{i:03d}', '10.00')

        self.assertEqual(len(self._lignes_du_classeur('&page_size=20')), 101)


class MultiDeviseTests(_SetupVentes):
    """
    ┌────────────────────────────────────────────────────────────────────────┐
    │ ON NE SOMME JAMAIS ENTRE DEVISES, ET ON LE DIT SUR LE DOCUMENT.        │
    │                                                                        │
    │ `ReportSpec` ne porte qu'UNE devise. Laisser les totaux de colonne      │
    │ s'additionner à travers deux monnaies produirait un nombre qui n'existe │
    │ pas, en bas d'une page que le marchand imprime et signe.               │
    └────────────────────────────────────────────────────────────────────────┘
    """

    def test_une_seule_devise_garde_ses_totaux_de_colonne(self):
        self._vente('VT-1', '100.00')
        spec = build_sales_report(
            Sale.objects.filter(organization=self.org), self.org,
            currency=self.org.currency,
        )
        self.assertIn('total', spec.group_totals)
        self.assertNotIn('ATTENTION', spec.subtitle)

    def test_deux_devises_RETIRENT_les_totaux_et_AVERTISSENT(self):
        self._seconde_devise()
        self._vente('VT-CDF', '100.00')
        self._vente('VT-USD', '50.00', devise='USD')

        spec = build_sales_report(
            Sale.objects.filter(organization=self.org), self.org,
            currency=self.org.currency,
        )
        self.assertEqual(spec.group_totals, ())
        self.assertIn('ATTENTION', spec.subtitle)
        self.assertIn('USD', spec.subtitle)

    def test_la_synthese_est_VENTILEE_par_devise(self):
        self._seconde_devise()
        self._vente('VT-CDF', '100.00')
        self._vente('VT-USD', '50.00', devise='USD', du='20.00')

        spec = build_sales_report(
            Sale.objects.filter(organization=self.org), self.org,
            currency=self.org.currency,
        )
        libelles = [intitule for intitule, _ in spec.summary]
        self.assertIn(f'Total {self.org.currency}', libelles)
        self.assertIn('Total USD', libelles)
        # Le reste dû n'est écrit que là où il existe : « Reste dû CDF : 0 »
        # sous une colonne soldée est du bruit qui masque la vraie créance.
        self.assertIn('Reste dû USD', libelles)
        self.assertNotIn(f'Reste dû {self.org.currency}', libelles)

    def test_chaque_LIGNE_s_ecrit_dans_SA_devise(self):
        """
        ┌────────────────────────────────────────────────────────────────────┐
        │ 120,75 USD NE SORT PAS « 121 » PARCE QUE L'ÉTABLISSEMENT EST EN    │
        │ FRANCS.                                                            │
        │                                                                    │
        │ Le spec ne posait pas `currency_field` alors que ses lignes portent │
        │ leur code et qu'une colonne « Devise » est affichée : TOUS les      │
        │ montants prenaient les décimales du document. Dans un              │
        │ établissement tenu en CDF (zéro décimale), un montant en dollars    │
        │ perdait ses centimes, sur un journal que le marchand signe.         │
        └────────────────────────────────────────────────────────────────────┘
        """
        self._seconde_devise()
        self._vente('VT-CDF', '50000.00')
        self._vente('VT-USD', '120.75', devise='USD')

        spec = build_sales_report(
            Sale.objects.filter(organization=self.org), self.org,
            currency=self.org.currency,
        )
        self.assertEqual(spec.currency_field, 'currency')

        resp = self.client.get(
            '/api/v1/sales/export/?export_format=csv', **self._headers,
        )
        texte = resp.content.decode('utf-8')
        self.assertIn('120,75', texte)
        # Le CDF n'a pas de décimale : la ligne en francs ne doit pas en gagner
        # parce que sa voisine en a besoin.
        self.assertIn('50 000', texte)
        self.assertNotIn('50 000,00', texte)

    def test_une_seule_devise_est_NOMMEE_meme_si_ce_n_est_pas_la_principale(self):
        """
        Un journal qui ne porte qu'une monnaie n'a rien à ventiler.

        On nomme donc CETTE devise plutôt que celle de l'établissement : sans
        cela, l'en-tête écrivait « Montants en FC » au-dessus de lignes en
        dollars, et le total général sortait sans ses centimes.
        """
        self._seconde_devise()
        self._vente('VT-USD', '120.75', devise='USD')

        spec = build_sales_report(
            Sale.objects.filter(organization=self.org), self.org,
            currency=self.org.currency,
        )
        self.assertEqual(spec.currency, 'USD')
        self.assertIsNone(spec.currency_field)
        # Une seule devise : les totaux de colonne existent, et ils sont justes.
        self.assertIn('total', spec.group_totals)

        resp = self.client.get(
            '/api/v1/sales/export/?export_format=csv', **self._headers,
        )
        self.assertIn('120,75', resp.content.decode('utf-8'))

    def test_un_journal_vide_garde_la_devise_de_l_etablissement(self):
        spec = build_sales_report(
            Sale.objects.filter(organization=self.org), self.org,
            currency=self.org.currency,
        )
        self.assertEqual(spec.currency, self.org.currency)
        self.assertIsNone(spec.currency_field)
